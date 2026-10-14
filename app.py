"""
app.py – Giao diện web Streamlit cho hệ thống đối soát thu chi bộ.

Khởi chạy: streamlit run app.py
"""
import io
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

# Thêm thư mục gốc vào sys.path để import các module backend
sys.path.insert(0, str(Path(__file__).resolve().parent))

import config
from config import SHEET_QUY, SHEET_PHAT
from statement_reader import read_statement, read_members, strip_recipient_info
from matcher import classify_transaction, match_member, get_quy_months
from main import run_core, apply_manual_assignments, PERIOD_OPTIONS

# ─────────────────────────────────────────────────────────────────────────────
# Cấu hình trang
# ─────────────────────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Đối Soát Thu Chi Quỹ Chi Bộ",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ─────────────────────────────────────────────────────────────────────────────
# CSS tùy chỉnh
# ─────────────────────────────────────────────────────────────────────────────
st.markdown("""
<style>
/* Google Font */
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&display=swap');

html, body, [class*="css"] { font-family: 'Inter', sans-serif; }

/* Hero header */
.hero {
    background: linear-gradient(135deg, #1a237e 0%, #283593 50%, #3949ab 100%);
    border-radius: 16px;
    padding: 32px 40px;
    margin-bottom: 28px;
    color: white;
    box-shadow: 0 8px 32px rgba(26,35,126,0.3);
}
.hero h1 { font-size: 2rem; font-weight: 700; margin: 0 0 6px 0; letter-spacing: -0.5px; }
.hero p  { font-size: 1rem; opacity: 0.85; margin: 0; }

/* Metric cards */
.metric-row { display: flex; gap: 16px; margin-bottom: 24px; }
.metric-card {
    flex: 1;
    background: white;
    border-radius: 12px;
    padding: 20px 24px;
    box-shadow: 0 2px 12px rgba(0,0,0,0.08);
    border-left: 4px solid;
    transition: transform .15s;
}
.metric-card:hover { transform: translateY(-2px); box-shadow: 0 6px 20px rgba(0,0,0,0.12); }
.metric-card.blue   { border-color: #3949ab; }
.metric-card.green  { border-color: #2e7d32; }
.metric-card.red    { border-color: #c62828; }
.metric-card.orange { border-color: #e65100; }
.metric-card .label { font-size: .8rem; color: #666; font-weight: 500; text-transform: uppercase; letter-spacing: .5px; }
.metric-card .value { font-size: 1.9rem; font-weight: 700; margin-top: 4px; }
.metric-card.blue   .value { color: #3949ab; }
.metric-card.green  .value { color: #2e7d32; }
.metric-card.red    .value { color: #c62828; }
.metric-card.orange .value { color: #e65100; }

/* Upload boxes */
.upload-section {
    background: #f8f9ff;
    border: 2px dashed #c5cae9;
    border-radius: 12px;
    padding: 20px;
    margin-bottom: 12px;
}

/* Status badges */
.badge {
    display: inline-block;
    padding: 2px 10px;
    border-radius: 12px;
    font-size: .78rem;
    font-weight: 600;
}
.badge-success { background: #e8f5e9; color: #2e7d32; }
.badge-danger  { background: #ffebee; color: #c62828; }
.badge-warning { background: #fff8e1; color: #e65100; }
.badge-info    { background: #e3f2fd; color: #1565c0; }

/* Tab styling */
.stTabs [data-baseweb="tab-list"] { gap: 8px; }
.stTabs [data-baseweb="tab"] {
    border-radius: 8px 8px 0 0;
    font-weight: 500;
    padding: 8px 20px;
}

/* Run button */
# div.stButton > button {
#     background: linear-gradient(135deg, #1a237e, #3949ab);
#     color: white;
#     border: none;
#     border-radius: 10px;
#     padding: 12px 32px;
#     font-size: 1rem;
#     font-weight: 600;
#     letter-spacing: .3px;
#     width: 100%;
#     transition: all .2s;
# }
# div.stButton > button:hover {
#     background: linear-gradient(135deg, #283593, #5c6bc0);
#     box-shadow: 0 4px 16px rgba(26,35,126,0.35);
#     transform: translateY(-1px);
# }

div.stButton > button {
    background: linear-gradient(135deg, #283593 0%, #3949ab 55%, #5c6bc0 100%) !important;
    color: #ffffff !important;
    border: 1px solid rgba(255,255,255,0.08) !important;
    box-shadow: 0 6px 18px rgba(57,73,171,0.28);
}

div.stButton > button:disabled,
div.stButton > button[disabled] {
    background: #e8eaf6 !important;
    color: #3949ab !important;
    border: 1px solid #c5cae9 !important;
    opacity: 1 !important;
}


/* Download button */
div[data-testid="stDownloadButton"] > button {
    background: #2e7d32;
    color: white;
    border: none;
    border-radius: 10px;
    padding: 10px 28px;
    font-weight: 600;
    width: 100%;
}
div[data-testid="stDownloadButton"] > button:hover {
    background: #1b5e20;
}

/* Section title */
.section-title {
    font-size: 1.05rem;
    font-weight: 700;
    color: #1a237e;
    margin: 20px 0 12px 0;
    padding-bottom: 6px;
    border-bottom: 2px solid #e8eaf6;
}

/* Sidebar */
section[data-testid="stSidebar"] > div {
    background: linear-gradient(180deg, #1a237e 0%, #283593 100%);
    padding-top: 16px;
}
section[data-testid="stSidebar"] * { color: white !important; }
section[data-testid="stSidebar"] .stFileUploader label { color: #c5cae9 !important; }

# /* Remove white background from file uploader in sidebar */
# section[data-testid="stSidebar"] [data-testid="stFileUploader"] section {
#     background-color: transparent !important;
#     border: 1px dashed rgba(255, 255, 255, 0.4) !important;
# }
# section[data-testid="stSidebar"] [data-testid="stFileUploader"] section > div {
#     background-color: transparent !important;
# }
# section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] {
#     background-color: transparent !important;
# }
# section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] > div {
#     background-color: transparent !important;
# }
# section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] button {
#     font-size: 0 !important;
#     color: transparent !important;
#     background-color: #ffffff !important;
#     border: 1px solid #c5cae9 !important;
#     border-radius: 8px !important;
#     display: inline-flex !important;
#     align-items: center !important;
#     justify-content: center !important;
#     padding: 0.5rem 1rem !important;
#     transition: all 0.2s;
# }
# section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] button * {
#     display: none !important;
# }
# section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] button::after {
#     content: "Chọn file";
#     font-size: 14px !important;
#     color: #1a237e !important;
#     font-weight: 600 !important;
#     display: block !important;
# }
# section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] button:hover {
#     background-color: #e8eaf6 !important;
#     border-color: #3949ab !important;
# }


# /* Translate Drag and drop text to Vietnamese */
# section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] > div > div > span {
#     display: none;
# }
# section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] > div > div::before {
#     content: "Kéo và thả file vào đây";
#     color: white;
#     font-size: 14px;
#     display: block;
#     margin-bottom: 2px;
# }
# section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] > div > div > small {
#     display: none;
# }
# section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] > div > div::after {
#     content: "Giới hạn 200MB/file • XLSX, XLS";
#     color: #c5cae9;
#     font-size: 12px;
#     display: block;
# }

/* Tối ưu lại phần File Uploader trong Sidebar */
section[data-testid="stSidebar"] [data-testid="stFileUploader"] {
    width: 100%;
}

section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] {
    background-color: rgba(255, 255, 255, 0.05) !important;
    border: 1px dashed rgba(255, 255, 255, 0.3) !important;
    border-radius: 8px !important;
    padding: 15px !important;
}

/* Ẩn hoàn toàn các thành phần mặc định gây đè chữ */
section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] button {
    display: none !important;
}

/* Ẩn text mặc định "Drag and drop file here" */
section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] > div > div > span {
    display: none !important;
}

/* Tạo giao diện nút bấm mới sạch sẽ */
section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"]::before {
    content: "Chọn file từ máy";
    display: block;
    background-color: white;
    color: #1a237e;
    text-align: center;
    padding: 8px 12px;
    border-radius: 6px;
    font-weight: 600;
    font-size: 14px;
    cursor: pointer;
    margin-bottom: 10px;
}

/* Cấu trúc lại phần text hướng dẫn bên dưới nút */
section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"]::after {
    content: "Kéo thả hoặc nhấn để chọn (Tối đa 200MB)";
    display: block;
    color: #c5cae9;
    font-size: 11px;
    text-align: center;
}

/* Hero gọn hơn */
.hero { padding: 24px 32px !important; margin-bottom: 20px !important; }
.hero h1 { font-size: 1.6rem !important; }
.metric-card .sub { font-size: .78rem; color: #888; margin-top: 2px; }

/* Các bước hướng dẫn (màn hình trống) */
.steps { display: grid; grid-template-columns: repeat(4, 1fr); gap: 14px; margin: 8px 0 20px 0; }
.step-card {
    background: #f8f9ff; border: 1px solid #e8eaf6; border-radius: 12px;
    padding: 16px 18px; display: flex; flex-direction: column; gap: 4px;
}
.step-card .n {
    width: 28px; height: 28px; border-radius: 50%; background: #3949ab; color: #fff;
    display: flex; align-items: center; justify-content: center; font-weight: 700; font-size: .9rem;
}
.step-card b { color: #1a237e; margin-top: 6px; }
.step-card span { color: #666; font-size: .85rem; }
@media (max-width: 900px) { .steps { grid-template-columns: repeat(2, 1fr); } }
</style>
""", unsafe_allow_html=True)
    

# ─────────────────────────────────────────────────────────────────────────────
# Helper
# ─────────────────────────────────────────────────────────────────────────────
def color_log_type(val: str) -> str:
    mapping = {
        "quy":     "background-color:#e3f2fd; color:#1565c0",
        "phat":    "background-color:#ffebee; color:#c62828",
        "dangphi": "background-color:#e8f5e9; color:#2e7d32",
        "ignore":  "background-color:#f3e5f5; color:#6a1b9a",
        "unknown": "background-color:#fff8e1; color:#e65100",
    }
    return mapping.get(str(val).lower(), "")


# ─────────────────────────────────────────────────────────────────────────────
# Tab quản lý mức đóng quỹ / mức phạt
# ─────────────────────────────────────────────────────────────────────────────
def _to_int(val) -> int:
    """Chuyển ô data_editor (có thể None/NaN) sang int, ô trống → 0."""
    return 0 if val is None or pd.isna(val) else int(val)


def render_settings_tab() -> None:
    """Giao diện chỉnh mức đóng quỹ và mức phạt, lưu vào settings.json."""
    # Tăng version để reset trạng thái các bảng sau khi lưu/khôi phục mặc định
    ver = st.session_state.setdefault("settings_ver", 0)
    settings = config.load_settings()
    rules = settings["phat_rules"]
    rule_labels = {r["col"]: r["label"] for r in rules}

    if msg := st.session_state.pop("settings_msg", None):
        st.success(msg)

    # ── Quỹ / Đảng phí ──────────────────────────────────────────────────────
    st.markdown('<div class="section-title">Mức đóng quỹ / Đảng phí</div>', unsafe_allow_html=True)
    c1, c2 = st.columns([1, 2])
    with c1:
        month_amount = st.number_input(
            "Mức đóng 1 tháng (đ)",
            min_value=0, step=1_000,
            value=int(settings["quy_month_amount"]),
            key=f"quy_month_{ver}",
        )
    with c2:
        st.markdown(
            f"Tự động nhận diện: **{month_amount:,}đ** = 1 tháng · "
            f"**{month_amount * 3:,}đ** = 1 quý · **{month_amount * 12:,}đ** = cả năm"
        )

    st.markdown("**Các mức khác cũng tính là đóng quỹ** (VD: chuyển dư phí ngân hàng)")
    extra_df = pd.DataFrame(
        settings["quy_extra_amount_hints"] or [], columns=["amount", "months"]
    ).rename(columns={"amount": "Số tiền (đ)", "months": "Số tháng"})
    extra_edit = st.data_editor(
        extra_df,
        num_rows="dynamic",
        width="stretch",
        key=f"quy_extra_{ver}",
        column_config={
            "Số tiền (đ)": st.column_config.NumberColumn(min_value=0, step=1_000, format="%d"),
            "Số tháng": st.column_config.SelectboxColumn(
                options=[1, 3, 12], help="1 = 1 tháng, 3 = 1 quý, 12 = cả năm"
            ),
        },
    )

    # ── Đảng phí ────────────────────────────────────────────────────────────
    st.markdown('<div class="section-title">Thu đảng phí</div>', unsafe_allow_html=True)
    d1, d2 = st.columns([1, 2])
    with d1:
        dang_phi = st.number_input(
            "Mức thu đảng phí (đ)", min_value=0, step=1_000,
            value=int(settings["dang_phi_amount"]), key=f"dang_phi_{ver}",
        )
    with d2:
        st.caption(
            "Giao dịch đúng số tiền này được phân loại là thu đảng phí, chỉ hiển thị trên giao diện "
            "và không cập nhật vào file báo cáo. Đặt 0 để tắt."
        )

    # ── Phạt ────────────────────────────────────────────────────────────────
    st.markdown('<div class="section-title">Mức phạt</div>', unsafe_allow_html=True)
    st.caption(
        "Nhận diện tối đa = N: giao dịch bằng Mức phạt × 1, 2, …, N được tự hiểu là khoản phạt này "
        "(0 = chỉ nhận diện theo từ khóa). Từ khóa phân cách bằng dấu phẩy, có dấu hay không đều được."
    )
    phat_df = pd.DataFrame([{
        "Cột": r["col"],
        "Loại vi phạm": r["label"],
        "Mức phạt (đ)": int(r.get("amount") or 0),
        "Nhận diện tối đa": int(r.get("max_multiple") or 0),
        "Từ khóa": ", ".join(r["keywords"]),
    } for r in rules])
    phat_edit = st.data_editor(
        phat_df,
        num_rows="fixed",
        hide_index=True,
        width="stretch",
        key=f"phat_rules_{ver}",
        column_config={
            "Cột": st.column_config.NumberColumn(disabled=True, width="small",
                                                  help="Cột tương ứng trong sheet phạt"),
            "Mức phạt (đ)": st.column_config.NumberColumn(min_value=0, step=1_000, format="%d"),
            "Nhận diện tối đa": st.column_config.NumberColumn(min_value=0, max_value=20, step=1),
            "Từ khóa": st.column_config.TextColumn(width="large"),
        },
    )

    unspecified_opts = [None] + [r["col"] for r in rules]
    cur_unspecified = settings["quy_unspecified_phat_col"]
    unspecified_col = st.selectbox(
        f"Giao dịch đúng {month_amount:,}đ (mức 1 tháng) nhưng không ghi rõ tháng/quý thì coi là",
        unspecified_opts,
        index=unspecified_opts.index(cur_unspecified) if cur_unspecified in unspecified_opts else 0,
        format_func=lambda c: "Quỹ của tháng giao dịch" if c is None else f"Phạt: {rule_labels[c]}",
        key=f"quy_unspec_{ver}",
    )

    # ── Gom cài đặt mới từ giao diện ────────────────────────────────────────
    new_settings = {
        "quy_month_amount": int(month_amount),
        "dang_phi_amount": int(dang_phi),
        "quy_extra_amount_hints": [
            {"amount": _to_int(r["Số tiền (đ)"]), "months": _to_int(r["Số tháng"])}
            for _, r in extra_edit.iterrows()
            if _to_int(r["Số tiền (đ)"]) > 0 and _to_int(r["Số tháng"]) in (1, 3, 12)
        ],
        "phat_rules": [{
            **rule,
            "label": str(r["Loại vi phạm"] or "").strip() or rule["label"],
            "amount": _to_int(r["Mức phạt (đ)"]),
            "max_multiple": _to_int(r["Nhận diện tối đa"]),
            "keywords": [k.strip() for k in str(r["Từ khóa"] or "").split(",") if k.strip()],
        } for rule, (_, r) in zip(rules, phat_edit.iterrows())],
        "quy_unspecified_phat_col": unspecified_col,
    }

    # ── Xem trước bảng nhận diện theo số tiền ───────────────────────────────
    quy_hints  = config.build_quy_amount_hints(new_settings["quy_month_amount"],
                                               new_settings["quy_extra_amount_hints"])
    phat_hints = config.build_phat_amount_hints(new_settings["phat_rules"])
    new_labels = {r["col"]: r["label"] for r in new_settings["phat_rules"]}
    period_text = {1: "Quỹ 1 tháng", 3: "Quỹ 1 quý", 12: "Quỹ cả năm"}

    with st.expander("Xem trước: số tiền nào sẽ được nhận diện thành khoản gì"):
        preview = [{"Số tiền": a, "Nhận diện là": period_text.get(m, f"Quỹ {m} tháng")}
                   for a, m in quy_hints.items()]
        preview += [{"Số tiền": a, "Nhận diện là": f"Phạt: {new_labels[c]}"}
                    for a, c in phat_hints.items() if a not in quy_hints]
        if dang_phi > 0:
            # Thu đảng phí được xét trước quỹ/phạt theo số tiền
            preview = [p for p in preview if p["Số tiền"] != dang_phi]
            preview.append({"Số tiền": dang_phi, "Nhận diện là": "Thu đảng phí (không ghi báo cáo)"})
        st.dataframe(
            pd.DataFrame(preview).sort_values("Số tiền").style.format({"Số tiền": "{:,.0f}đ"}),
            hide_index=True, width="stretch",
        )
        st.caption("Giao dịch có từ khóa (phạt/quỹ) trong nội dung luôn được ưu tiên hơn số tiền.")

    clashes = sorted(a for a in phat_hints if a in quy_hints)
    if clashes:
        st.warning(
            "Các mức sau trùng giữa quỹ và phạt, sẽ được tính là **quỹ** nếu nội dung không có từ khóa phạt: "
            + ", ".join(f"{a:,}đ" for a in clashes)
        )

    errors = []
    if new_settings["quy_month_amount"] <= 0:
        errors.append("Mức đóng 1 tháng phải lớn hơn 0.")
    for r in new_settings["phat_rules"]:
        if r["max_multiple"] > 0 and r["amount"] <= 0:
            errors.append(f"“{r['label']}”: cần nhập Mức phạt để nhận diện theo số tiền.")
    for e in errors:
        st.error(e)

    # ── Lưu / khôi phục ─────────────────────────────────────────────────────
    b1, b2, _ = st.columns([1, 1, 2])
    with b1:
        if st.button("Lưu cài đặt", disabled=bool(errors), width="stretch"):
            config.save_settings(new_settings)
            _after_settings_change("Đã lưu cài đặt mức thu mới.")
    with b2:
        if st.button("Khôi phục mặc định", width="stretch"):
            config.save_settings(config.default_settings())
            _after_settings_change("Đã khôi phục mức thu mặc định.")


def _after_settings_change(msg: str) -> None:
    """Xoá kết quả đối soát cũ (tính theo mức cũ) và tải lại giao diện."""
    had_result = st.session_state.pop("result", None) is not None
    if had_result:
        msg += " Kết quả đối soát trước đó đã được xoá – vui lòng chạy lại đối soát."
    st.session_state["settings_msg"] = msg
    st.session_state["settings_ver"] += 1
    st.rerun()


# ─────────────────────────────────────────────────────────────────────────────
# Kiểm tra nhanh file tải lên (cache theo nội dung file)
# ─────────────────────────────────────────────────────────────────────────────
@st.cache_data(show_spinner=False)
def inspect_statement(data: bytes) -> dict:
    try:
        df = read_statement(io.BytesIO(data))
    except Exception as e:
        return {"error": str(e)}
    if df.empty:
        return {"error": "Không có giao dịch phát sinh Có nào trong sao kê."}
    return {
        "n": len(df), "total": int(df["amount"].sum()),
        "d0": df["date"].min().strftime("%d/%m/%Y"), "d1": df["date"].max().strftime("%d/%m/%Y"),
    }


@st.cache_data(show_spinner=False)
def inspect_report(data: bytes) -> dict:
    try:
        wb = load_workbook(io.BytesIO(data))
        missing = [n for n in (SHEET_QUY, SHEET_PHAT) if n not in wb.sheetnames]
        if missing:
            return {"error": f"Thiếu sheet: {', '.join(missing)}"}
        return {
            "n_quy": len(read_members(wb[SHEET_QUY])),
            "n_phat": len(read_members(wb[SHEET_PHAT])),
        }
    except Exception as e:
        return {"error": str(e)}


# ─────────────────────────────────────────────────────────────────────────────
# Sidebar – Upload files
# ─────────────────────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown("## Quản lý quỹ CBSV1")
    st.markdown("---")
    st.markdown("### Tải lên file dữ liệu")

    st.markdown("**1. File sao kê ngân hàng**")
    sao_ke_file = st.file_uploader(
        "Sao kê (Excel)", type=["xlsx", "xls"], key="sao_ke", label_visibility="collapsed",
    )
    sk_info = inspect_statement(sao_ke_file.getvalue()) if sao_ke_file else None
    if sk_info and "error" in sk_info:
        st.error(f"Sao kê không hợp lệ: {sk_info['error']}")
    elif sk_info:
        st.success(f"✓ {sao_ke_file.name}")
        st.caption(
            f"{sk_info['n']} giao dịch có · {sk_info['total']:,}đ · "
            f"{sk_info['d0']} → {sk_info['d1']}"
        )

    st.markdown("**2. File báo cáo Đảng phí**")
    tong_hop_file = st.file_uploader(
        "Báo cáo (Excel)", type=["xlsx", "xls"], key="tong_hop", label_visibility="collapsed",
    )
    bc_info = inspect_report(tong_hop_file.getvalue()) if tong_hop_file else None
    if bc_info and "error" in bc_info:
        st.error(f"Báo cáo không hợp lệ: {bc_info['error']}")
    elif bc_info:
        st.success(f"✓ {tong_hop_file.name}")
        st.caption(f"{bc_info['n_quy']} đảng viên (sheet quỹ) · {bc_info['n_phat']} (sheet phạt)")

    st.markdown("---")
    

files_ok = (
    sao_ke_file is not None and tong_hop_file is not None
    and "error" not in sk_info and "error" not in bc_info
)


# ─────────────────────────────────────────────────────────────────────────────
# Tab "Cần xử lý": gán thủ công các giao dịch app chưa nhận diện/khớp tên được
# ─────────────────────────────────────────────────────────────────────────────
KIND_LABEL = {"Quỹ": "quy", "Phạt": "phat", "Đảng phí": "dangphi", "Bỏ qua": "ignore"}


def _period_label(months: list) -> str:
    if len(months) == 12:
        return "Cả năm"
    if len(months) == 3:
        return f"Quý {(months[0] - 1) // 3 + 1}"
    return f"Tháng {months[0]}"


def _build_manual_df(result: dict, todo: pd.DataFrame) -> pd.DataFrame:
    """Điền sẵn gợi ý (loại khoản thu, đảng viên, lỗi vi phạm, kỳ đóng) cho từng dòng."""
    phat_label_of = {r["col"]: r["label"] for r in config.PHAT_RULES}
    rows = []
    for idx, log in todo.iterrows():
        tx = {"details": log["Nội dung"], "amount": int(log["Số tiền"]),
              "date": datetime.strptime(log["Ngày"], "%d/%m/%Y")}
        tx_type, col, _ = classify_transaction(tx)
        member, _, _ = match_member(
            {"search_text": strip_recipient_info(tx["details"])}, result["members_quy"]
        )
        rows.append({
            "_idx": idx,
            "Ngày": log["Ngày"],
            "Số tiền": tx["amount"],
            "Nội dung": log["Nội dung"],
            "Lý do chưa xử lý": log["Trạng thái"],
            "Xử lý như": {"quy": "Quỹ", "phat": "Phạt"}.get(tx_type),
            "Đảng viên": member["name"] if member else None,
            "Loại vi phạm": phat_label_of.get(col),
            "Kỳ đóng quỹ": _period_label(get_quy_months(tx)),
        })
    return pd.DataFrame(rows)


def render_manual_tab(result: dict, todo: pd.DataFrame) -> None:
    if todo.empty:
        st.success("Không còn giao dịch nào cần xử lý thủ công.")
        return

    st.markdown('<div class="section-title">Giao dịch cần bạn xác nhận</div>', unsafe_allow_html=True)
    st.caption(
        "App đã điền sẵn gợi ý. Với mỗi dòng, chọn **Xử lý như** (Quỹ / Phạt / Bỏ qua), "
        "đảng viên và (nếu là phạt) loại vi phạm, rồi bấm **Áp dụng**. "
        "Dòng để trống ô *Xử lý như* sẽ được giữ nguyên."
    )

    ver = st.session_state.setdefault("manual_ver", 0)
    cache_key = f"manual_df_{ver}"
    if cache_key not in st.session_state:
        st.session_state[cache_key] = _build_manual_df(result, todo)
    base_df = st.session_state[cache_key]

    names = sorted({m["name"] for m in result["members_quy"]})
    edited = st.data_editor(
        base_df,
        hide_index=True,
        width="stretch",
        num_rows="fixed",
        height=min(80 + 36 * len(base_df), 520),
        key=f"manual_editor_{ver}",
        column_config={
            "_idx": None,
            "Ngày": st.column_config.TextColumn(disabled=True, width="small"),
            "Số tiền": st.column_config.NumberColumn(disabled=True, format="%d đ", width="small"),
            "Nội dung": st.column_config.TextColumn(disabled=True, width="large"),
            "Lý do chưa xử lý": st.column_config.TextColumn(disabled=True),
            "Xử lý như": st.column_config.SelectboxColumn(options=list(KIND_LABEL), width="small"),
            "Đảng viên": st.column_config.SelectboxColumn(options=names, width="medium"),
            "Loại vi phạm": st.column_config.SelectboxColumn(
                options=[r["label"] for r in config.PHAT_RULES], width="medium"),
            "Kỳ đóng quỹ": st.column_config.SelectboxColumn(options=PERIOD_OPTIONS, width="small"),
        },
    )

    n_ready = int(edited["Xử lý như"].notna().sum())
    if st.button(f"Áp dụng {n_ready} dòng đã chọn", disabled=n_ready == 0, type="primary"):
        col_of = {r["label"]: r["col"] for r in config.PHAT_RULES}
        assignments = [{
            "idx": int(r["_idx"]),
            "kind": KIND_LABEL[r["Xử lý như"]],
            "member": r["Đảng viên"],
            "phat_col": col_of.get(r["Loại vi phạm"]),
            "period": r["Kỳ đóng quỹ"],
        } for _, r in edited.iterrows() if r["Xử lý như"] in KIND_LABEL]
        applied, errors = apply_manual_assignments(result, assignments)
        st.session_state["manual_ver"] = ver + 1
        st.session_state["manual_msg"] = (applied, errors)
        st.rerun()


# ─────────────────────────────────────────────────────────────────────────────
# Main area
# ─────────────────────────────────────────────────────────────────────────────
tab_doisoat, tab_settings = st.tabs(["Đối soát", "Quản lý mức thu"])

with tab_settings:
    render_settings_tab()

with tab_doisoat:
    if msg := st.session_state.pop("manual_msg", None):
        applied, errors = msg
        if applied:
            st.success(f"Đã áp dụng {applied} dòng vào báo cáo.")
        for e in errors:
            st.warning(e)

    if "result" not in st.session_state:
        st.markdown("""
        <div class="steps">
            <div class="step-card"><div class="n">1</div><b>Tải sao kê</b><span>File Excel sao kê ngân hàng ở thanh bên trái</span></div>
            <div class="step-card"><div class="n">2</div><b>Tải báo cáo</b><span>File báo cáo đảng phí cần cập nhật</span></div>
            <div class="step-card"><div class="n">3</div><b>Chạy đối soát</b><span>App tự phân loại & điền vào báo cáo</span></div>
            <div class="step-card"><div class="n">4</div><b>Xác nhận & tải về</b><span>Xử lý dòng còn sót, rồi tải file đã cập nhật</span></div>
        </div>
        """, unsafe_allow_html=True)

    col_run, col_hint = st.columns([1, 2], vertical_alignment="center")
    with col_run:
        run_clicked = st.button(
            "Chạy đối soát", disabled=not files_ok, width="stretch", type="primary",
        )
    with col_hint:
        if not files_ok:
            st.caption("Cần tải lên đủ 2 file hợp lệ ở thanh bên trái để bắt đầu.")
        elif "result" in st.session_state:
            st.caption("Chạy lại sẽ làm mới kết quả và xoá các chỉnh sửa thủ công chưa tải về.")

    if run_clicked and files_ok:
        with st.spinner("Đang xử lý dữ liệu..."):
            try:
                result = run_core(
                    io.BytesIO(sao_ke_file.getvalue()),
                    io.BytesIO(tong_hop_file.getvalue()),
                )
                st.session_state["result"]   = result
                st.session_state["out_name"] = f"{Path(tong_hop_file.name).stem}_UPDATED.xlsx"
                st.session_state["manual_ver"] = st.session_state.get("manual_ver", 0) + 1
                st.toast("Đối soát hoàn tất!", icon="")
            except Exception as e:
                st.error(f"Lỗi khi xử lý: {e}")
                st.exception(e)

    if "result" in st.session_state:
        result  = st.session_state["result"]
        logs_df = result["logs"]

        # ── Thống kê tổng quan ──────────────────────────────────────────────
        is_old = logs_df["Trạng thái"].str.startswith("Bỏ qua vì số bút toán")
        needs_action = (
            logs_df["Loại"].isin(["quy", "phat", "unknown"]) & (logs_df["Tên khớp"] == "") & ~is_old
        )
        todo      = logs_df[needs_action]
        written   = logs_df[(logs_df["Tên khớp"] != "") & logs_df["Loại"].isin(["quy", "phat"])]
        n_matched = len(written)
        total_quy  = written.loc[written["Loại"] == "quy", "Số tiền"].sum()
        total_phat = written.loc[written["Loại"] == "phat", "Số tiền"].sum()
        total_dp   = logs_df.loc[logs_df["Loại"] == "dangphi", "Số tiền"].sum()

        st.markdown(f"""
        <div class="metric-row">
            <div class="metric-card blue">
                <div class="label">Tổng giao dịch</div>
                <div class="value">{result['n_transactions']}</div>
                <div class="sub">{int(is_old.sum())} dòng đã xử lý từ lần trước</div>
            </div>
            <div class="metric-card green">
                <div class="label">Đã ghi vào báo cáo</div>
                <div class="value">{n_matched}</div>
                <div class="sub">đã khớp được đảng viên</div>
            </div>
            <div class="metric-card red">
                <div class="label">Cần xử lý</div>
                <div class="value">{len(todo)}</div>
                <div class="sub">xem tab Cần xử lý</div>
            </div>
            <div class="metric-card orange">
                <div class="label">Tổng thu</div>
                <div class="value">{total_quy + total_phat + total_dp:,.0f}đ</div>
                <div class="sub">Quỹ {total_quy:,.0f}đ · Phạt {total_phat:,.0f}đ · Đảng phí {total_dp:,.0f}đ</div>
            </div>
        </div>
        """, unsafe_allow_html=True)

        tab_log, tab_todo, tab_dl = st.tabs([
            "Log giao dịch",
            f"Cần xử lý ({len(todo)})",
            "Tải xuống",
        ])

        # ── Log giao dịch ───────────────────────────────────────────────────
        with tab_log:
            st.markdown('<div class="section-title">Nhật ký đối soát giao dịch</div>', unsafe_allow_html=True)
            c1, c2, c3 = st.columns([3, 2, 2])
            with c1:
                search = st.text_input("Tìm theo tên hoặc nội dung", placeholder="Nhập từ khoá...")
            with c2:
                filter_type = st.selectbox("Lọc loại", ["Tất cả", "quy", "phat", "dangphi", "ignore", "unknown"])
            with c3:
                filter_matched = st.selectbox("Khớp tên", ["Tất cả", "Đã khớp", "Chưa khớp"])
            only_todo = st.checkbox("Chỉ hiện các dòng cần xử lý")

            df_show = logs_df.copy()
            if search:
                mask = (
                    df_show["Tên khớp"].str.contains(search, case=False, na=False)
                    | df_show["Nội dung"].str.contains(search, case=False, na=False)
                )
                df_show = df_show[mask]
            if filter_type != "Tất cả":
                df_show = df_show[df_show["Loại"] == filter_type]
            if filter_matched == "Đã khớp":
                df_show = df_show[df_show["Tên khớp"] != ""]
            elif filter_matched == "Chưa khớp":
                df_show = df_show[df_show["Tên khớp"] == ""]
            if only_todo:
                df_show = df_show[needs_action.reindex(df_show.index)]

            st.caption(f"Hiển thị {len(df_show)} / {len(logs_df)} giao dịch")
            styled = (
                df_show.reset_index(drop=True)
                .style
                .map(color_log_type, subset=["Loại"])
                .format({"Số tiền": "{:,.0f}", "Điểm": lambda x: f"{x:.1f}" if pd.notna(x) and x != "" else ""})
            )
            st.dataframe(styled, width="stretch", height=420)
            st.download_button(
                "Tải bảng đang hiển thị (CSV)",
                data=df_show.to_csv(index=False).encode("utf-8-sig"),
                file_name="DoiSoat_Log.csv", mime="text/csv",
            )

        with tab_todo:
            render_manual_tab(result, todo)

        # ── Tải xuống ───────────────────────────────────────────────────────
        with tab_dl:
            st.markdown('<div class="section-title">Tải xuống file Excel đã cập nhật</div>', unsafe_allow_html=True)
            out_name = st.session_state.get("out_name", "BaoCao_UPDATED.xlsx")
            if len(todo):
                st.warning(f"Còn {len(todo)} giao dịch chưa được ghi vào báo cáo (xem tab Cần xử lý).")

            col_dl, col_info = st.columns([1, 2])
            with col_dl:
                st.download_button(
                    label="Tải file Excel đã cập nhật",
                    data=result["excel_bytes"],
                    file_name=out_name,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    width="stretch", type="primary",
                )
            with col_info:
                st.markdown(f"""
                **File:** `{out_name}`

                File này bao gồm:
                - Sheet **{SHEET_QUY}** – đã cập nhật trạng thái nộp quỹ
                - Sheet **{SHEET_PHAT}** – đã cập nhật các khoản phạt
                - Sheet **DoiSoat_Log** – nhật ký toàn bộ giao dịch đối soát
                """)
